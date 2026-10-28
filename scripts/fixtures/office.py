"""Independent standards-complete Office fixtures; no product OOXML builders used."""
from pathlib import Path
import sys
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

def create(directory):
 directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
 doc=Document();doc.add_heading('SmaltCode 引擎往返样本',0);doc.add_paragraph('中文内容与原件保护。');doc.save(directory/'sample.docx')
 wb=Workbook();ws=wb.active;ws.title='成本';ws['A1']=12;ws['B1']='=A1*2';wb.save(directory/'sample.xlsx')
 deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[0]);slide.shapes.title.text='SmaltCode 引擎样本';deck.save(directory/'sample.pptx')
if __name__=='__main__':create(sys.argv[1])
