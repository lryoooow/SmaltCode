from pathlib import Path
import sys,json
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
p=Path(sys.argv[1]);p.mkdir(parents=True,exist_ok=True)
if len(sys.argv)==2:
 w=Workbook();s=w.active;s.title='成本'
 for row in [['项目','金额'],['甲',20],['乙',10],['丙',30]]:s.append(row)
 s['D1']='=SUM(B2:B4)';s['D2']='=IFERROR(1/0,7)';s['D3']='=1/0';s['D4']='=DATE(2026,10,27)';s['D4'].number_format='yyyy-mm-dd'
 t=w.create_sheet('汇总');t['A1']="='成本'!D1*2"
 dv=DataValidation(type='whole',operator='between',formula1=0,formula2=100);s.add_data_validation(dv);dv.add(s['B2']);s['A1'].comment=Comment('NATIVE_COMMENT_PRESERVED','fixture')
 chart=BarChart();chart.add_data(Reference(s,min_col=2,min_row=1,max_row=4),titles_from_data=True);s.add_chart(chart,'G2');w.save(p/'source.xlsx')
else:
 w=load_workbook(p/'working.xlsx');s=w['成本'];assert len(s._charts)==1;assert len(s.data_validations.dataValidation)==1;assert s['A1'].comment.text=='NATIVE_COMMENT_PRESERVED';assert s['E1'].value=='=literal';assert s['E2'].value.startswith('=');assert w['新增']['A1'].value==123
 v=load_workbook(p/'working.xlsx',data_only=True);assert v['成本']['E2'].value==60;assert v['汇总']['A1'].value==120
 print(json.dumps({'charts':1,'validation':1,'nativeComment':True,'recalculated':True}))
