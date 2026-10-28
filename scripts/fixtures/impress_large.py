from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Inches
p=Path(sys.argv[1]);p.mkdir(parents=True,exist_ok=True);r=Presentation()
for i in range(500):
 s=r.slides.add_slide(r.slide_layouts[6]);s.name=('大型演示页面_'+str(i)+'_')*5
 if i==0:
  for j in range(150):s.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1)).text=('中文与明确引用"'+str(j))*250
r.save(p/'working.pptx')
