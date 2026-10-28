import sys,zipfile,json
from pathlib import Path
from pptx import Presentation
from pypdf import PdfReader
p=Path(sys.argv[1]);r=Presentation(p/'working.pptx');assert len(r.slides)==2
s=r.slides[0];assert any(getattr(a,'text','')=='已核对的中文标题'for a in s.shapes);assert any(a.shape_type==13 for a in s.shapes);assert 'ENGINE_NATIVE_NOTES'in s.notes_slide.notes_text_frame.text
with zipfile.ZipFile(p/'working.pptx')as z:assert any(b'ENGINE_NATIVE_COMMENT'in z.read(n)for n in z.namelist()if 'comment'in n.lower())
pdf=PdfReader(p/'export.pdf');assert len(pdf.pages)==2;assert any(page.extract_text().strip()for page in pdf.pages);print(json.dumps({'slides':2,'image':True,'notes':True,'comment':True,'pdfPages':2}))
