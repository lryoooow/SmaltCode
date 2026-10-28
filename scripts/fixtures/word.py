"""Independent complex Word fixture: Chinese typography, tables, images, fields and sections."""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Inches,Pt,Cm
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw

def create(directory):
 d=Path(directory);d.mkdir(parents=True,exist_ok=True)
 image=Image.new('RGB',(420,150),'#eef2ff');draw=ImageDraw.Draw(image);draw.rectangle((20,20,400,130),outline='#6366f1',width=4);draw.text((40,60),'SmaltCode Word',fill='#1e293b');image.save(d/'sample.png')
 doc=Document();style=doc.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(10);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 section=doc.sections[0];section.top_margin=Cm(2);section.bottom_margin=Cm(2);section.header.paragraphs[0].text='HEADER_PRESERVED · SmaltCode 中文文档'
 footer=section.footer.paragraphs[0];footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页 · FOOTER_PRESERVED')
 doc.add_heading('SmaltCode Word 回归样本',0);doc.add_paragraph('AI_SELECTION_TARGET');p=doc.add_paragraph();p.add_run('中文粗体 ').bold=True;p.add_run('italic example ').italic=True;p.add_run('多样式段落必须保留。')
 doc.add_picture(str(d/'sample.png'),width=Inches(3))
 for i in range(3):doc.add_paragraph('列表项目 '+str(i+1),style='List Bullet')
 table=doc.add_table(rows=1,cols=3);table.style='Table Grid'
 for c,t in zip(table.rows[0].cells,['序号','中文说明','数值']):c.text=t
 for i in range(75):
  cells=table.add_row().cells
  for c,t in zip(cells,[str(i+1),'跨页表格保留内容 '+str(i+1),str((i+1)*3)]):c.text=t
 doc.add_section(WD_SECTION_START.NEW_PAGE);doc.add_heading('第二节',1);doc.add_paragraph('SECTION_TWO_PRESERVED · 中文分页与页码。');doc.save(d/'rich.docx')
if __name__=='__main__':create(sys.argv[1])
