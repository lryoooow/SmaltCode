"""Independent verification of actual Word OOXML and PDF rendering, after native engine roundtrip."""
from pathlib import Path
import json,sys,zipfile,io,hashlib
from docx import Document
from PIL import Image,ImageChops,ImageStat
import pypdfium2 as pdfium
from pypdf import PdfReader

def verify(directory):
 directory=Path(directory);source=Document(directory/'rich.docx');saved=Document(directory/'working.docx');assert len(saved.sections)>=len(source.sections);assert len(saved.tables)>=2;assert any('跨页表格' in c.text for row in saved.tables[0].rows for c in row.cells);assert any('ENGINE_TABLE' in c.text for table in saved.tables for row in table.rows for c in row.cells)
 assert 'HEADER_PRESERVED' in '\n'.join(p.text for s in saved.sections for p in s.header.paragraphs)
 assert 'FOOTER_PRESERVED' in '\n'.join(p.text for s in saved.sections for p in s.footer.paragraphs)
 with zipfile.ZipFile(directory/'working.docx')as z:
  document=z.read('word/document.xml').decode();comments=z.read('word/comments.xml').decode();assert 'ENGINE_NATIVE_COMMENT' in comments;assert '<w:ins' in document;assert 'ENGINE_TRACKED_INSERTION' in document;assert 'Microsoft YaHei' in document or 'Microsoft YaHei' in z.read('word/styles.xml').decode()
 before=PdfReader(directory/'before.pdf');after=PdfReader(directory/'after.pdf');assert len(before.pages)>=3;assert len(after.pages)>=len(before.pages);assert 'HEADER_PRESERVED' in after.pages[0].extract_text();assert any('ENGINE_TABLE' in p.extract_text()for p in after.pages)
 images=[]
 for name in ['before.pdf','after.pdf']:
  pdf=pdfium.PdfDocument(str(directory/name));page=pdf[0];images.append(page.render(scale=1.5).to_pil().convert('RGB'));page.close();pdf.close()
 assert images[0].size==images[1].size
 diffs={};width,height=images[0].size
 for label,box in [('header',(0,0,width,72)),('footer',(0,height-67,width,height))]:
  clips=[im.crop(box)for im in images];mean=sum(ImageStat.Stat(ImageChops.difference(*clips)).mean)/3;diffs[label]=mean;assert mean<1.0,(label,mean)
 images[1].save(directory/'page-one.png')
 result={'independentDocx':True,'pdfPagesBefore':len(before.pages),'pdfPagesAfter':len(after.pages),'headerFooterPixelMeanError':diffs,'nativeComment':True,'nativeRevisions':True,'tables':len(saved.tables),'sections':len(saved.sections),'pageOne':str(directory/'page-one.png')};(directory/'independent-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':verify(sys.argv[1])
