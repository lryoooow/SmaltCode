"""Reusable isolated artifact desktop harness."""
from _desktop_control import publish
import json,os,socket,subprocess,sys,threading,time
from datetime import datetime,timezone
from pathlib import Path
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
root=Path(__file__).resolve().parents[1];sys.path.insert(0,str(root/'.qa-python'))
from playwright.sync_api import sync_playwright,expect
expect.set_options(timeout=30000)
run=root/'.qa-work'/('artifact-ui-'+str(time.time_ns()));run.mkdir(parents=True)
shots=root/'docs/screenshots/phase-nine';shots.mkdir(parents=True,exist_ok=True)
requests=[];release=threading.Event();held=0;errors=[];checks=[];instances=[]
class Model(BaseHTTPRequestHandler):
 def log_message(self,*args):pass
 def do_GET(self):
  self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(json.dumps({'data':[{'id':'remote-model','context_window':256000,'supported_parameters':['tools']}]}).encode())
 def do_POST(self):
  global held
  body=json.loads(self.rfile.read(int(self.headers.get('Content-Length',0))));requests.append(body);messages=body['messages'];latest=max(i for i,m in enumerate(messages)if m['role']=='user');prompt=str(messages[latest]['content']);done=any(m['role']=='tool'for m in messages[latest+1:]);call=None
  self.send_response(200);self.send_header('Content-Type','text/event-stream');self.end_headers()
  try:
   if 'HOLD_REMOTE'in prompt and not done:
    held+=1;self.wfile.write(('data: '+json.dumps({'choices':[{'delta':{'reasoning_content':'检查远程项目中的实际文件。'},'finish_reason':None}]})+'\n\n').encode());self.wfile.flush();release.wait(50)
   if not done and 'SSH_SETUP_ROUTE'in prompt:call=('shell',{'program':'ssh','args':['-o','BatchMode=yes','user@127.0.0.1','hostname']})
   if not done and 'SSH_DISCOVER_ROUTE'in prompt:call=('ssh_connections',{})
   if not done and 'ASK_PHONE'in prompt:call=('ask_user',{'question':'请确认需要的输出语言'})
   if not done and 'GENERATE_DELIVERABLE'in prompt:call=('resource_write',{'root':'output','path':'最终报告.md','content':'# 独立产物\nDELIVERABLE_ONLY_OUTPUT\n'})
   if not done and 'EDIT_HOST'in prompt:call=('write',{'path':'notes.md','content':'HOST_CHANGED_BY_REMOTE\n'})
   delta={'tool_calls':[{'index':0,'id':'call_'+str(len(requests)),'function':{'name':call[0],'arguments':json.dumps(call[1],ensure_ascii=False)}}]}if call else{'content':'已完成远程检查。运行位置：配对主机。项目测试尚未运行。'}
   self.wfile.write(('data: '+json.dumps({'choices':[{'delta':delta,'finish_reason':'tool_calls'if call else'stop'}]},ensure_ascii=False)+'\n\ndata: [DONE]\n\n').encode());self.wfile.flush()
  except(BrokenPipeError,ConnectionResetError):pass
server=ThreadingHTTPServer(('127.0.0.1',0),Model);threading.Thread(target=server.serve_forever,daemon=True).start()
def until(fn,seconds=35):
 end=time.time()+seconds
 while time.time()<end:
  if fn():return
  time.sleep(.1)
 raise AssertionError('condition timed out')
class Desktop:
 def __init__(self,label):
  self.dir=run/label;self.home=self.dir/'home';self.workspace=self.dir/'project';(self.home/'threads').mkdir(parents=True);self.workspace.mkdir();(self.workspace/'notes.md').write_text(label+'_FILE\n',encoding='utf8');self.label=label;self.process=None;self.browser=None;self.page=None
  with socket.socket()as sock:sock.bind(('127.0.0.1',0));self.port=sock.getsockname()[1]
  self.quit=self.dir/'quit';self.capture=self.dir/'capture';self.zoom=self.dir/'zoom';self.ack=self.dir/'ack';self.deny=self.dir/'deny';self.export=self.dir/'export-target';self.active=self.dir/'active-window'
  at=datetime.now(timezone.utc).isoformat()
  for ident,title in [('th_remote_shared','远程项目对话'),('th_remote_other','独立浏览对话')]:
   meta={'id':ident,'title':title,'cwd':str(self.workspace),'model':'remote-model','provider':'openai','permissionMode':'default','createdAt':at,'updatedAt':at,'turnCount':1,'lastOutcome':'completed'}
   events=[{'t':'meta','meta':meta},{'t':'turn-started','turnId':'seed','at':at}]
   for i in range(120 if ident=='th_remote_shared'else 1):events.extend([{'t':'user-msg','id':'u-'+str(i),'text':'历史问题 '+str(i),'at':at},{'t':'agent-msg','id':'a-'+str(i),'text':'可安全查看的项目回复 '+str(i),'at':at}])
   events.append({'t':'turn-ended','turnId':'seed','at':at});(self.home/'threads'/(ident+'.jsonl')).write_text('\n'.join(json.dumps(v,ensure_ascii=False)for v in events)+'\n',encoding='utf8')
  (self.home/'workspace-access.json').write_text(json.dumps([str(self.workspace.resolve()).lower()]),encoding='utf8')
  (self.home/'settings.json').write_text(json.dumps({'defaultProvider':'openai','defaultModel':'remote-model','providers':{'openai':{'baseUrl':f'http://127.0.0.1:{server.server_port}/v1','apiKey':'HOST_MODEL_KEY_NEVER_EXPOSE'}}}),encoding='utf8')
  paths={k:str(getattr(self,k))for k in ['quit','capture','zoom','ack','deny','export','active']};self.harness=self.dir/'harness.cjs';self.harness.write_text("const {app,dialog,BrowserWindow}=require('electron'),fs=require('node:fs');app.getVersion=()=>"+json.dumps(json.loads((root/'packages/desktop/package.json').read_text(encoding='utf8'))['version'])+";const paths="+json.dumps(paths)+";dialog.showMessageBox=async()=>({response:fs.existsSync(paths.deny)?1:0});dialog.showSaveDialog=async()=>fs.existsSync(paths.export)?({canceled:false,filePath:fs.readFileSync(paths.export,'utf8')}):({canceled:true});dialog.showOpenDialog=async()=>({canceled:false,filePaths:["+json.dumps(str(self.workspace))+"]});setInterval(()=>{if(fs.existsSync(paths.quit))app.quit();const windows=BrowserWindow.getAllWindows(),selected=fs.existsSync(paths.active)?fs.readFileSync(paths.active,'utf8'):'';const w=windows.find(w=>w.webContents.getTitle()===selected)||windows[0];if(w&&fs.existsSync(paths.capture)){const file=fs.readFileSync(paths.capture,'utf8');fs.unlinkSync(paths.capture);w.webContents.capturePage().then(image=>fs.writeFileSync(file,image.toPNG()))}if(w&&fs.existsSync(paths.zoom)){const value=Number(fs.readFileSync(paths.zoom,'utf8'));if(Number.isFinite(value)&&value>0){w.webContents.setZoomFactor(value);fs.unlinkSync(paths.zoom);fs.writeFileSync(paths.ack,'done')}}},100).unref();require("+json.dumps(str(root/'packages/desktop/dist/main/index.js'))+");",encoding='utf8')
  self.log=open(self.dir/'electron.log','w',encoding='utf8');instances.append(self)
 def launch(self,p):
  self.quit.unlink(missing_ok=True);env={**os.environ,'SMALT_HOME':str(self.home)}
  for key in list(env):
   if any(v in key.upper()for v in ['API_KEY','TOKEN','SECRET','PASSWORD','CREDENTIAL'])or key=='ELECTRON_RUN_AS_NODE':env.pop(key,None)
  self.process=subprocess.Popen([str(getattr(self,'runtime_executable',os.environ.get('SMALT_TEST_ELECTRON_EXECUTABLE',root/'node_modules/electron/dist/electron.exe'))),str(self.harness),'--development-preview','--identity-test','--remote-debugging-port='+str(self.port)],env=env,stdout=self.log,stderr=self.log,creationflags=subprocess.CREATE_NO_WINDOW)
  for _ in range(150):
   try:self.browser=p.chromium.connect_over_cdp(f'http://127.0.0.1:{self.port}',timeout=500);break
   except Exception:time.sleep(.1)
  assert self.browser;until(lambda:any(v.url.startswith('file:')for c in self.browser.contexts for v in c.pages));self.page=next(v for c in self.browser.contexts for v in c.pages if v.url.startswith('file:'));self.page.on('pageerror',lambda e:errors.append(self.label+':'+str(e)));self.page.set_default_timeout(30000);self.page.wait_for_load_state('networkidle');self.page.evaluate("cwd=>{localStorage.setItem('smalt.ui.onboarded','true');localStorage.setItem('smalt.ui.dark','false');localStorage.setItem('smalt.ui.workspace',JSON.stringify(cwd))}",str(self.workspace));self.page.reload();self.page.wait_for_load_state('networkidle')
  if not getattr(self,'local_default_only',False):
   self.nav('打开资料与产物');self.page.get_by_role('tab',name='执行环境',exact=True).click();self.nav('添加可修改的工作目录');expect(self.page.get_by_test_id('file-execution-path')).to_have_text(str(self.workspace));self.nav('关闭资料与产物')
 def api(self,action,value=None):return self.page.evaluate('v=>window.smalt.remote(v.action,v.value)',{'action':action,'value':value or{}})
 def state(self):return self.api('list')
 def nav(self,label):self.page.get_by_role('button',name=label,exact=True).click()
 def shot(self,name,page=None):
  publish(self.active,page.title()if page else self.page.title(),encoding='utf8')
  target=shots/name;target.unlink(missing_ok=True);publish(self.capture,str(target),encoding='utf8');until(lambda:target.exists())
 def close(self):
  if self.process and self.process.poll()is None:self.quit.write_text('quit');self.process.wait(timeout=20);assert self.process.returncode==0

def make_office(app):
 import zipfile
 def zipdoc(name,parts):
  with zipfile.ZipFile(app.workspace/name,'w',zipfile.ZIP_DEFLATED)as archive:
   for path,value in parts.items():archive.writestr(path,value)
 zipdoc('项目说明.docx',{'[Content_Types].xml':'<Types/>','word/document.xml':'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:rPr><w:b/></w:rPr><w:t>项目说明</w:t></w:r></w:p><w:p><w:r><w:t>原始项目结论，等待核对。</w:t></w:r></w:p></w:body></w:document>','word/header1.xml':'<w:hdr xmlns:w="urn:w">HEADER_UNCHANGED</w:hdr>','word/styles.xml':'<styles>STYLE_UNCHANGED</styles>'})
 zipdoc('项目数据.xlsx',{'[Content_Types].xml':'<Types/>','xl/workbook.xml':'<workbook xmlns:r="urn:r"><sheets><sheet name="成本" r:id="r1"/></sheets></workbook>','xl/_rels/workbook.xml.rels':'<Relationships><Relationship Id="r1" Target="worksheets/sheet1.xml"/></Relationships>','xl/worksheets/sheet1.xml':'<worksheet><sheetData><row r="1"><c r="A1" s="1"><v>12</v></c><c r="B1"><f>A1*2</f><v>24</v></c><c r="C1" t="inlineStr"><is><t>备注</t></is></c></row></sheetData></worksheet>','xl/styles.xml':'<styleSheet><cellXfs><xf numFmtId="0"/><xf numFmtId="0"/></cellXfs></styleSheet>'})
 # Create standards-complete Office fixtures with independent libraries too.
 from docx import Document
 from docx.shared import Pt
 doc=Document();doc.add_heading('项目说明',level=1);doc.add_paragraph('原始项目结论，等待核对。');doc.sections[0].header.paragraphs[0].text='HEADER_UNCHANGED';doc.save(str(app.workspace/'项目说明.docx'))
 from openpyxl import Workbook
 from openpyxl.styles import Font
 wb=Workbook();ws=wb.active;ws.title='成本';ws['A1']=12;ws['A1'].font=Font(bold=True);ws['B1']='=A1*2';ws['C1']='备注';other=wb.create_sheet('进度');other['A1']='保留另一张表';wb.save(str(app.workspace/'项目数据.xlsx'))
 from reportlab.pdfgen import canvas
 output=canvas.Canvas(str(app.workspace/'审阅说明.pdf'),pagesize=(400,500))
 for i in range(2):output.drawString(30,450,'Review document page '+str(i+1));output.showPage()
 output.save()
def review(app,name):
 app.page.evaluate('(name)=>window.smalt.artifacts("open",{threadId:"th_remote_shared",path:name,dark:false})',name)
 until(lambda:any(v.url.endswith('/artifact.html')and not v.is_closed()and v.title()==name+' · SmaltCode 审阅'for c in app.browser.contexts for v in c.pages))
 result=next(v for c in app.browser.contexts for v in c.pages if v.url.endswith('/artifact.html')and not v.is_closed()and v.title()==name+' · SmaltCode 审阅');result.on('pageerror',lambda e:errors.append(str(e)));result.set_default_timeout(30000);return result
def state(page):return page.evaluate('()=>window.smaltArtifact.invoke("view")')
def button(page,label):page.get_by_role('button',name=label,exact=True).click()
